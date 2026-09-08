import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

print("Creating Banque Chaabi / Banque Populaire PFA PowerPoint presentation...")

prs = Presentation()
prs.slide_width = Inches(13.333)  # 16:9 widescreen
prs.slide_height = Inches(7.5)

# Color Palette: Banque Chaabi / Banque Populaire
ORANGE_CHAABI = RGBColor(243, 111, 33)   # #F36F21
MARRON_CHAABI = RGBColor(61, 12, 2)      # #3D0C02
WHITE = RGBColor(255, 255, 255)
DARK_TEXT = RGBColor(31, 41, 55)         # #1F2937
GRAY_BG = RGBColor(250, 250, 250)        # #FAFAFA
LIGHT_ORANGE = RGBColor(255, 244, 237)   # #FFF4ED
BORDER_ORANGE = RGBColor(255, 219, 194)  # #FFDBC2

blank_slide_layout = prs.slide_layouts[6]

# Utility function to add background color
def set_bg(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

# Utility function to add card/shape
def add_card(slide, left, top, width, height, bg_color, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape

# ==========================================
# SLIDE 1: COUVERTURE (TITLE SLIDE)
# ==========================================
slide1 = prs.slides.add_slide(blank_slide_layout)
set_bg(slide1, ORANGE_CHAABI)

# Big top-right white circle for Banque Chaabi Logo
circle_top = slide1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.5), Inches(-1.5), Inches(7.5), Inches(7.5))
circle_top.fill.solid()
circle_top.fill.fore_color.rgb = WHITE
circle_top.line.fill.background()

# Big bottom-right white circle for EMSI Logo
circle_bot = slide1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.5), Inches(3.2), Inches(5.5), Inches(5.5))
circle_bot.fill.solid()
circle_bot.fill.fore_color.rgb = WHITE
circle_bot.line.fill.background()

# Insert Banque Chaabi Logo inside top circle
bp_logo_path = 'frontend/public/logo_BP.png'
if os.path.exists(bp_logo_path):
    slide1.shapes.add_picture(bp_logo_path, Inches(8.2), Inches(1.2), width=Inches(4.2))

# Insert EMSI Logo inside bottom circle
emsi_logo_path = 'scratch/emsi_logo.png'
if os.path.exists(emsi_logo_path):
    slide1.shapes.add_picture(emsi_logo_path, Inches(9.2), Inches(4.8), width=Inches(3.8))

# Title Text Box on Left
tx_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(6.5), Inches(5.0))
tf = tx_box.text_frame
tf.word_wrap = True

p1 = tf.paragraphs[0]
p1.text = "PFA : Plateforme Web de Data Lineage & Supervision Bancaire"
p1.font.bold = True
p1.font.size = Pt(36)
p1.font.color.rgb = WHITE

p2 = tf.add_paragraph()
p2.text = "\nRéalisé et présenté par :"
p2.font.size = Pt(20)
p2.font.bold = True
p2.font.color.rgb = MARRON_CHAABI

p3 = tf.add_paragraph()
p3.text = "Mlle. Meriam Erragragui\nMr. Omar Tazi"
p3.font.size = Pt(22)
p3.font.bold = True
p3.font.color.rgb = WHITE

p4 = tf.add_paragraph()
p4.text = "\nEncadrée par : Pr. Chaimae Taoussi"
p4.font.size = Pt(18)
p4.font.color.rgb = MARRON_CHAABI

p5 = tf.add_paragraph()
p5.text = "Entreprise d'accueil : Banque Populaire (Banque Chaabi)\nAnnée Universitaire 2024/2025"
p5.font.size = Pt(16)
p5.font.color.rgb = WHITE

# ==========================================
# SLIDE 2: SOMMAIRE (AGENDA)
# ==========================================
slide2 = prs.slides.add_slide(blank_slide_layout)
set_bg(slide2, WHITE)

# Title: Sommaire
title_box = slide2.shapes.add_textbox(Inches(4.5), Inches(0.5), Inches(4.3), Inches(1.0))
tf2 = title_box.text_frame
p = tf2.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run1 = p.add_run()
run1.text = "Som"
run1.font.bold = True
run1.font.size = Pt(40)
run1.font.color.rgb = MARRON_CHAABI

run2 = p.add_run()
run2.text = "maire"
run2.font.bold = True
run2.font.size = Pt(40)
run2.font.color.rgb = ORANGE_CHAABI

# 6 Grid Cards for Sommaire
sommaire_items = [
    ("1. Introduction", ORANGE_CHAABI),
    ("2. Cadre du projet", MARRON_CHAABI),
    ("3. Conception", RGBColor(224, 92, 14)),
    ("4. Technologies utilisées", RGBColor(139, 58, 43)),
    ("5. Implémentation", RGBColor(217, 83, 30)),
    ("6. Conclusion", ORANGE_CHAABI)
]

coords = [
    (Inches(1.2), Inches(2.0)),
    (Inches(5.0), Inches(2.0)),
    (Inches(8.8), Inches(2.0)),
    (Inches(1.2), Inches(4.5)),
    (Inches(5.0), Inches(4.5)),
    (Inches(8.8), Inches(4.5))
]

for idx, (item_text, color) in enumerate(sommaire_items):
    left, top = coords[idx]
    card = add_card(slide2, left, top, Inches(3.3), Inches(1.8), color)
    tf_c = card.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.alignment = PP_ALIGN.CENTER
    p_c.text = item_text
    p_c.font.bold = True
    p_c.font.size = Pt(22)
    p_c.font.color.rgb = WHITE

# ==========================================
# SLIDE 3: INTRODUCTION (CONTEXT & PROBLEMATIC)
# ==========================================
slide3 = prs.slides.add_slide(blank_slide_layout)
set_bg(slide3, WHITE)

# Left Side Banner
banner = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(4.2), Inches(7.5))
banner.fill.solid()
banner.fill.fore_color.rgb = ORANGE_CHAABI
banner.line.fill.background()

tb_intro = slide3.shapes.add_textbox(Inches(0.4), Inches(3.0), Inches(3.4), Inches(2.0))
tf_intro = tb_intro.text_frame
p_i = tf_intro.paragraphs[0]
r_i1 = p_i.add_run()
r_i1.text = "Intro"
r_i1.font.bold = True
r_i1.font.size = Pt(38)
r_i1.font.color.rgb = MARRON_CHAABI

r_i2 = p_i.add_run()
r_i2.text = "duction"
r_i2.font.bold = True
r_i2.font.size = Pt(38)
r_i2.font.color.rgb = WHITE

# Right Content Box
content_box3 = slide3.shapes.add_textbox(Inches(4.8), Inches(0.8), Inches(7.8), Inches(6.0))
tf3 = content_box3.text_frame
tf3.word_wrap = True

p_c1 = tf3.paragraphs[0]
p_c1.text = "➢ Contexte"
p_c1.font.bold = True
p_c1.font.size = Pt(26)
p_c1.font.color.rgb = MARRON_CHAABI

bullets_ctx = [
    "Besoin d'outils numériques modernes pour automatiser la gouvernance des données bancaires.",
    "Complexité croissante des flux entre comptes, agences, transactions et rapports mensuels.",
    "Exigence stricte d'audit de conformité (PII, Data Owner, traçabilité des transformations)."
]
for b in bullets_ctx:
    p_b = tf3.add_paragraph()
    p_b.text = f"  ▪  {b}"
    p_b.font.size = Pt(18)
    p_b.font.color.rgb = DARK_TEXT

p_c2 = tf3.add_paragraph()
p_c2.text = "\n➢ Problématique"
p_c2.font.bold = True
p_c2.font.size = Pt(26)
p_c2.font.color.rgb = ORANGE_CHAABI

p_prob = tf3.add_paragraph()
p_prob.text = "Comment concevoir une plateforme simple, intelligente et évolutive qui aide la Banque Populaire à :"
p_prob.font.bold = True
p_prob.font.size = Pt(18)
p_prob.font.color.rgb = DARK_TEXT

bullets_prob = [
    "Cartographier automatiquement les 6 tables métier et leur Column-Level Lineage.",
    "Identifier la classification PII et les propriétaires de données (Data Owners).",
    "Interroger le catalogue en langage naturel et vocal via un Assistant IA intelligent (RAG Mistral)."
]
for b in bullets_prob:
    p_b = tf3.add_paragraph()
    p_b.text = f"  ▪  {b}"
    p_b.font.size = Pt(18)
    p_b.font.color.rgb = DARK_TEXT

# ==========================================
# SLIDE 4: CADRE DU PROJET
# ==========================================
slide4 = prs.slides.add_slide(blank_slide_layout)
set_bg(slide4, WHITE)

# Left Square Card
card4 = add_card(slide4, Inches(0.8), Inches(1.2), Inches(4.5), Inches(5.0), ORANGE_CHAABI)
tf_c4 = card4.text_frame
tf_c4.word_wrap = True
p_c4 = tf_c4.paragraphs[0]
p_c4.alignment = PP_ALIGN.CENTER
r_c4_1 = p_c4.add_run()
r_c4_1.text = "Cadre "
r_c4_1.font.bold = True
r_c4_1.font.size = Pt(36)
r_c4_1.font.color.rgb = MARRON_CHAABI

r_c4_2 = p_c4.add_run()
r_c4_2.text = "du projet"
r_c4_2.font.bold = True
r_c4_2.font.size = Pt(36)
r_c4_2.font.color.rgb = WHITE

# Right Details Box
details_box4 = slide4.shapes.add_textbox(Inches(5.8), Inches(1.2), Inches(6.8), Inches(5.0))
tf4 = details_box4.text_frame
tf4.word_wrap = True

p_head4 = tf4.paragraphs[0]
p_head4.text = "Projet réalisé dans le cadre d'un Projet de Fin d'Année (PFA)\n"
p_head4.font.bold = True
p_head4.font.size = Pt(22)
p_head4.font.color.rgb = MARRON_CHAABI

items4 = [
    ("Thème : ", "Automatisation & Investigation Intelligente du Data Lineage Bancaire"),
    ("Entreprise d'accueil : ", "Banque Populaire (Banque Chaabi)"),
    ("Durée : ", "2 mois"),
    ("Encadrement : ", "École EMSI (Pr. Chaimae Taoussi)")
]

for label, val in items4:
    p_i = tf4.add_paragraph()
    r_l = p_i.add_run()
    r_l.text = f"▪  {label}"
    r_l.font.bold = True
    r_l.font.size = Pt(20)
    r_l.font.color.rgb = ORANGE_CHAABI

    r_v = p_i.add_run()
    r_v.text = val
    r_v.font.size = Pt(20)
    r_v.font.color.rgb = DARK_TEXT

# ==========================================
# SLIDE 5: CONCEPTION (SECTION TRANSITION)
# ==========================================
slide5 = prs.slides.add_slide(blank_slide_layout)
set_bg(slide5, ORANGE_CHAABI)

# Big Center White Circle
circle_sec = slide5.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.4), Inches(0.5), Inches(6.5), Inches(6.5))
circle_sec.fill.solid()
circle_sec.fill.fore_color.rgb = WHITE
circle_sec.line.fill.background()

tf5 = circle_sec.text_frame
tf5.word_wrap = True
p5_sec = tf5.paragraphs[0]
p5_sec.alignment = PP_ALIGN.CENTER
r5_1 = p5_sec.add_run()
r5_1.text = "\n\nConcep"
r5_1.font.bold = True
r5_1.font.size = Pt(48)
r5_1.font.color.rgb = MARRON_CHAABI

r5_2 = p5_sec.add_run()
r5_2.text = "tion"
r5_2.font.bold = True
r5_2.font.size = Pt(48)
r5_2.font.color.rgb = ORANGE_CHAABI

# ==========================================
# SLIDE 6: TECHNOLOGIES UTILISÉES
# ==========================================
slide6 = prs.slides.add_slide(blank_slide_layout)
set_bg(slide6, WHITE)

# Title: Technologies
title_box6 = slide6.shapes.add_textbox(Inches(4.0), Inches(0.5), Inches(5.3), Inches(1.0))
tf6 = title_box6.text_frame
p6 = tf6.paragraphs[0]
p6.alignment = PP_ALIGN.CENTER
r6_1 = p6.add_run()
r6_1.text = "Tech"
r6_1.font.bold = True
r6_1.font.size = Pt(40)
r6_1.font.color.rgb = MARRON_CHAABI

r6_2 = p6.add_run()
r6_2.text = "nologies"
r6_2.font.bold = True
r6_2.font.size = Pt(40)
r6_2.font.color.rgb = ORANGE_CHAABI

tech_boxes = [
    ("Frontend Web", "React.js, Vite, Lucide Icons\nWeb Speech API (Speech-to-Text & Text-to-Speech)\nDesign System Banque Populaire", ORANGE_CHAABI),
    ("Backend API & BDD", "Python, FastAPI, SQLAlchemy\nPostgreSQL (10 254 transactions, 324 alertes, comptes, clients)", MARRON_CHAABI),
    ("Gouvernance & Lineage", "Apache Atlas REST API\nColumn-Level Lineage & Metadata Catalog\nAnalyse d'impact dynamique", RGBColor(224, 92, 14)),
    ("IA & Vector Database", "Mistral LLM Engine\nQdrant Vector Database\nMoteur RAG d'investigation en langage naturel", RGBColor(139, 58, 43))
]

coords_tech = [
    (Inches(1.2), Inches(1.8)),
    (Inches(7.0), Inches(1.8)),
    (Inches(1.2), Inches(4.5)),
    (Inches(7.0), Inches(4.5))
]

for idx, (head_t, body_t, color_t) in enumerate(tech_boxes):
    l, t = coords_tech[idx]
    card = add_card(slide6, l, t, Inches(5.2), Inches(2.2), LIGHT_ORANGE, BORDER_ORANGE)
    tf_tc = card.text_frame
    tf_tc.word_wrap = True
    
    p_h = tf_tc.paragraphs[0]
    p_h.text = f"⚡ {head_t}"
    p_h.font.bold = True
    p_h.font.size = Pt(20)
    p_h.font.color.rgb = color_t
    
    for line in body_t.split('\n'):
        p_b = tf_tc.add_paragraph()
        p_b.text = f"▪ {line}"
        p_b.font.size = Pt(15)
        p_b.font.color.rgb = DARK_TEXT

# ==========================================
# SLIDE 7: IMPLÉMENTATION & DÉMONSTRATION
# ==========================================
slide7 = prs.slides.add_slide(blank_slide_layout)
set_bg(slide7, WHITE)

# Left Side Card
card7 = add_card(slide7, Inches(0.8), Inches(1.2), Inches(4.5), Inches(5.0), MARRON_CHAABI)
tf_c7 = card7.text_frame
tf_c7.word_wrap = True
p_c7 = tf_c7.paragraphs[0]
p_c7.alignment = PP_ALIGN.CENTER
r_c7_1 = p_c7.add_run()
r_c7_1.text = "Implémen"
r_c7_1.font.bold = True
r_c7_1.font.size = Pt(36)
r_c7_1.font.color.rgb = WHITE

r_c7_2 = p_c7.add_run()
r_c7_2.text = "tation"
r_c7_2.font.bold = True
r_c7_2.font.size = Pt(36)
r_c7_2.font.color.rgb = ORANGE_CHAABI

# Right Features Box
features_box7 = slide7.shapes.add_textbox(Inches(5.8), Inches(1.2), Inches(6.8), Inches(5.0))
tf7 = features_box7.text_frame
tf7.word_wrap = True

feat_items = [
    ("Catalogue de Données : ", "6 tables métier synchronisées avec Atlas (clients, agences, comptes, transactions, transactions_suspectes, reporting_mensuel)."),
    ("Analyse d'Impact Dinamique : ", "Identification instantanée des Data Owners (ex: Équipe Monétique) et classifications PII (Confidentiel / Sensible / Critique)."),
    ("Column-Level Lineage : ", "Visualisation claire du cheminement des colonnes dans les pipelines ETL (ex: detection_transactions_suspectes)."),
    ("Assistant IA Vocal : ", "Interaction vocale native (Speech-to-Text) et lecture audio à la demande (Text-to-Speech) connectée au RAG Mistral."),
    ("Supervision Temps Réel : ", "Dashboard interactif connecté directement aux tables PostgreSQL.")
]

for label, desc in feat_items:
    p_f = tf7.add_paragraph()
    r_l = p_f.add_run()
    r_l.text = f"▪ {label}"
    r_l.font.bold = True
    r_l.font.size = Pt(17)
    r_l.font.color.rgb = ORANGE_CHAABI

    r_d = p_f.add_run()
    r_d.text = desc
    r_d.font.size = Pt(17)
    r_d.font.color.rgb = DARK_TEXT

# ==========================================
# SLIDE 8: CONCLUSION
# ==========================================
slide8 = prs.slides.add_slide(blank_slide_layout)
set_bg(slide8, ORANGE_CHAABI)

circle_c = slide8.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.4), Inches(0.5), Inches(6.5), Inches(6.5))
circle_c.fill.solid()
circle_c.fill.fore_color.rgb = WHITE
circle_c.line.fill.background()

tf8 = circle_c.text_frame
tf8.word_wrap = True
p8 = tf8.paragraphs[0]
p8.alignment = PP_ALIGN.CENTER
r8_1 = p8.add_run()
r8_1.text = "\n\nConclu"
r8_1.font.bold = True
r8_1.font.size = Pt(44)
r8_1.font.color.rgb = MARRON_CHAABI

r8_2 = p8.add_run()
r8_2.text = "sion"
r8_2.font.bold = True
r8_2.font.size = Pt(44)
r8_2.font.color.rgb = ORANGE_CHAABI

p8_sub = tf8.add_paragraph()
p8_sub.alignment = PP_ALIGN.CENTER
p8_sub.text = "Merci pour votre attention !"
p8_sub.font.bold = True
p8_sub.font.size = Pt(24)
p8_sub.font.color.rgb = MARRON_CHAABI

output_pptx = "scratch/Presentation_Projet_Banking_Lineage_Banque_Chaabi.pptx"
prs.save(output_pptx)
print(f"POWERPOINT GENERATED SUCCESSFULLY AT: {output_pptx}")
